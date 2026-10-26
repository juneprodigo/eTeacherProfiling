import os
import io
import datetime
import openpyxl
from openpyxl.drawing.image import Image as OpenpyxlImage
from django.conf import settings

TEMPLATE_PATH = os.path.join(settings.BASE_DIR.parent, 'CS Form No. 212 Revised 2026 Personal Data Sheet PDS.xlsx')


class PDSExcelExporter:
    """
    Populates an employee's complete Personal Data Sheet (CS Form No. 212 Revised 2026)
    into the official Excel template with exact cell coordinates and formatting.
    """

    def __init__(self, teacher, template_path=None):
        self.teacher = teacher
        self.template_path = template_path or TEMPLATE_PATH
        if not os.path.exists(self.template_path):
            alt_path = os.path.join(os.path.dirname(os.path.dirname(os.path.abspath(__file__))), '..', 'CS Form No. 212 Revised 2026 Personal Data Sheet PDS.xlsx')
            if os.path.exists(alt_path):
                self.template_path = alt_path

        self.wb = openpyxl.load_workbook(self.template_path, data_only=False)

    def _format_date(self, val):
        if not val:
            return ""
        if isinstance(val, (datetime.date, datetime.datetime)):
            return val.strftime("%d/%m/%Y")
        try:
            d = datetime.date.fromisoformat(str(val).split('T')[0])
            return d.strftime("%d/%m/%Y")
        except Exception:
            return str(val)

    def _safe_str(self, val):
        if val is None or val == 'None':
            return ""
        return str(val).strip()

    def generate(self):
        self._populate_c1()
        self._populate_c2()
        self._populate_c3()
        self._populate_c4()
        if 'C1' in self.wb.sheetnames:
            self.wb.active = self.wb['C1']
        for sheet in self.wb.worksheets:
            try:
                sheet.sheet_view.topLeftCell = 'A1'
            except Exception:
                pass
        return self.wb

    def save_to_stream(self):
        self.generate()
        stream = io.BytesIO()
        self.wb.save(stream)
        stream.seek(0)
        return stream

    def _populate_c1(self):
        ws = self.wb['C1']
        t = self.teacher
        pi = getattr(t, 'personal_info', None)

        # I. PERSONAL INFORMATION
        surname = getattr(pi, 'surname', '') or (t.user.last_name if t.user else '')
        first_name = getattr(pi, 'first_name', '') or (t.user.first_name if t.user else '')
        middle_name = getattr(pi, 'middle_name', '')
        name_extension = getattr(pi, 'name_extension', '')

        ws['D10'] = self._safe_str(surname).upper()
        ws['D11'] = self._safe_str(first_name).upper()
        ws['N11'] = self._safe_str(name_extension).upper()
        ws['D12'] = self._safe_str(middle_name).upper()

        if pi and pi.date_of_birth:
            ws['D13'] = self._format_date(pi.date_of_birth)
        ws['D15'] = self._safe_str(getattr(pi, 'place_of_birth', ''))

        # Sex at Birth (D16 Male, E16 Female)
        sex = (getattr(pi, 'sex_at_birth', '') or '').strip().lower()
        if sex == 'male':
            ws['D16'] = True
            ws['E16'] = False
        elif sex == 'female':
            ws['D16'] = False
            ws['E16'] = True

        # Civil Status (D17 Single, E17 Married, D18 Widowed, D20 Separated, E19 Other)
        cs = (getattr(pi, 'civil_status', '') or '').strip().lower()
        ws['D17'] = (cs == 'single')
        ws['E17'] = (cs == 'married')
        ws['D18'] = ('widow' in cs)
        ws['D20'] = ('separated' in cs)
        ws['E19'] = ('other' in cs or 'solo' in cs)

        if pi and pi.height_m:
            ws['D22'] = float(pi.height_m)
        if pi and pi.weight_kg:
            ws['D24'] = float(pi.weight_kg)
        ws['D25'] = self._safe_str(getattr(pi, 'blood_type', ''))

        ws['D27'] = self._safe_str(getattr(pi, 'umid_id_no', ''))
        ws['D29'] = self._safe_str(getattr(pi, 'pag_ibig_id_no', ''))
        ws['D31'] = self._safe_str(getattr(pi, 'philhealth_no', ''))
        ws['D32'] = self._safe_str(getattr(pi, 'philsys_pcn', ''))
        ws['D33'] = self._safe_str(getattr(pi, 'tin_no', ''))
        ws['D34'] = self._safe_str(getattr(pi, 'agency_employee_no', '') or t.employee_id)

        # Citizenship
        is_filipino = getattr(pi, 'citizenship_filipino', True)
        is_dual = getattr(pi, 'dual_citizenship', False)
        ws['J13'] = bool(is_filipino)
        ws['K13'] = bool(is_dual)
        if is_dual:
            if getattr(pi, 'dual_citizenship_country', ''):
                ws['L15'] = self._safe_str(pi.dual_citizenship_country)
            ws['L14'] = bool(getattr(pi, 'dual_citizenship_by_birth', False))
            ws['M14'] = bool(getattr(pi, 'dual_citizenship_by_naturalization', False))

        # Residential Address
        ws['I17'] = self._safe_str(getattr(pi, 'res_house_block_lot', ''))
        ws['L17'] = self._safe_str(getattr(pi, 'res_street', ''))
        ws['I21'] = self._safe_str(getattr(pi, 'res_subdivision_village', ''))
        ws['L21'] = self._safe_str(getattr(pi, 'res_barangay', ''))
        ws['I23'] = self._safe_str(getattr(pi, 'res_city_municipality', ''))
        ws['L23'] = self._safe_str(getattr(pi, 'res_province', ''))
        ws['I24'] = self._safe_str(getattr(pi, 'res_zip_code', ''))

        # Permanent Address
        ws['I25'] = self._safe_str(getattr(pi, 'perm_house_block_lot', ''))
        ws['L25'] = self._safe_str(getattr(pi, 'perm_street', ''))
        ws['I28'] = self._safe_str(getattr(pi, 'perm_subdivision_village', ''))
        ws['L28'] = self._safe_str(getattr(pi, 'perm_barangay', ''))
        ws['I30'] = self._safe_str(getattr(pi, 'perm_city_municipality', ''))
        ws['L30'] = self._safe_str(getattr(pi, 'perm_province', ''))
        ws['I31'] = self._safe_str(getattr(pi, 'perm_zip_code', ''))

        ws['I32'] = self._safe_str(getattr(pi, 'telephone_no', ''))
        ws['I33'] = self._safe_str(getattr(pi, 'mobile_no', ''))
        ws['I34'] = self._safe_str(getattr(pi, 'email_address', '') or (t.user.email if t.user else ''))

        # II. FAMILY BACKGROUND
        fb = getattr(t, 'family_background', None)
        if fb:
            ws['D36'] = self._safe_str(fb.spouse_surname).upper()
            ws['D37'] = self._safe_str(fb.spouse_first_name).upper()
            ws['G37'] = self._safe_str(fb.spouse_name_extension).upper()
            ws['D38'] = self._safe_str(fb.spouse_middle_name).upper()
            ws['D39'] = self._safe_str(fb.spouse_occupation)
            ws['D40'] = self._safe_str(fb.spouse_employer_business_name)
            ws['D41'] = self._safe_str(fb.spouse_business_address)
            ws['D42'] = self._safe_str(fb.spouse_telephone_no)

            ws['D44'] = self._safe_str(fb.father_surname).upper()
            ws['D45'] = self._safe_str(fb.father_first_name).upper()
            ws['G45'] = self._safe_str(fb.father_name_extension).upper()
            ws['D46'] = self._safe_str(fb.father_middle_name).upper()

            ws['D48'] = self._safe_str(fb.mother_surname).upper()
            ws['D49'] = self._safe_str(fb.mother_first_name).upper()
            ws['D50'] = self._safe_str(fb.mother_middle_name).upper()

        # Children (rows 37 to 49 on C1)
        children = list(fb.children.all().order_by('date_of_birth', 'id')) if fb else []
        c7_sheet = self.wb['C7_Family Background cont. '] if 'C7_Family Background cont. ' in self.wb.sheetnames else None

        for idx, child in enumerate(children):
            c_dob = self._format_date(child.date_of_birth)
            if idx < 13:
                r = 37 + idx
                ws[f'I{r}'] = self._safe_str(child.name).upper()
                ws[f'M{r}'] = c_dob
            elif c7_sheet:
                cont_r = 7 + (idx - 13)
                if cont_r <= 50:
                    c7_sheet[f'B{cont_r}'] = self._safe_str(child.name).upper()
                    c7_sheet[f'G{cont_r}'] = c_dob

        if len(children) <= 13 and c7_sheet:
            c7_sheet.sheet_state = 'hidden'

        # III. EDUCATIONAL BACKGROUND
        edu_rows = {
            'ELEMENTARY': 55,
            'SECONDARY': 56,
            'VOCATIONAL/TRADE COURSE': 57,
            'VOCATIONAL': 57,
            'COLLEGE': 58,
            'GRADUATE STUDIES': 59,
            'GRADUATE': 59
        }
        c8_sheet = self.wb['C8_Educ. Background cont.'] if 'C8_Educ. Background cont.' in self.wb.sheetnames else None
        used_levels = set()
        c8_idx = 0

        for edu in t.educational_backgrounds.all().order_by('id'):
            lvl_key = (edu.level or '').strip().upper()
            target_row = None
            for k, r in edu_rows.items():
                if k in lvl_key:
                    if r not in used_levels:
                        target_row = r
                        used_levels.add(r)
                    break

            if target_row:
                ws[f'D{target_row}'] = self._safe_str(edu.school_name)
                ws[f'G{target_row}'] = self._safe_str(edu.degree_course)
                ws[f'J{target_row}'] = self._safe_str(edu.from_year)
                ws[f'K{target_row}'] = self._safe_str(edu.to_year)
                ws[f'L{target_row}'] = self._safe_str(edu.highest_level_units_earned)
                ws[f'M{target_row}'] = self._safe_str(edu.year_graduated)
                ws[f'N{target_row}'] = self._safe_str(edu.scholarship_academic_honors)
            elif c8_sheet:
                c8_r = 7 + c8_idx
                if c8_r <= 50:
                    c8_sheet[f'B{c8_r}'] = self._safe_str(edu.level)
                    c8_sheet[f'D{c8_r}'] = self._safe_str(edu.school_name)
                    c8_sheet[f'G{c8_r}'] = self._safe_str(edu.degree_course)
                    c8_sheet[f'J{c8_r}'] = self._safe_str(edu.from_year)
                    c8_sheet[f'K{c8_r}'] = self._safe_str(edu.to_year)
                    c8_sheet[f'L{c8_r}'] = self._safe_str(edu.highest_level_units_earned)
                    c8_sheet[f'M{c8_r}'] = self._safe_str(edu.year_graduated)
                    c8_sheet[f'N{c8_r}'] = self._safe_str(edu.scholarship_academic_honors)
                    c8_idx += 1

        if c8_idx == 0 and c8_sheet:
            c8_sheet.sheet_state = 'hidden'

    def _populate_c2(self):
        ws = self.wb['C2']
        t = self.teacher

        # IV. CIVIL SERVICE ELIGIBILITY (rows 5 to 11 on C2: A=Career Service, F=Rating, G=Date, I=Place, L=License No, M=Validity)
        c9_sheet = self.wb['C9_Elig cont.'] if 'C9_Elig cont.' in self.wb.sheetnames else None
        eligs = list(t.civil_service_eligibilities.all().order_by('id'))
        for idx, elig in enumerate(eligs):
            exam_date = self._format_date(elig.examination_date)
            val_date = self._format_date(elig.license_validity_date)
            if idx < 7:
                r = 5 + idx
                ws[f'A{r}'] = self._safe_str(elig.eligibility_type)
                ws[f'F{r}'] = self._safe_str(elig.rating)
                ws[f'G{r}'] = exam_date
                ws[f'I{r}'] = self._safe_str(elig.examination_place)
                ws[f'L{r}'] = self._safe_str(elig.license_number)
                ws[f'M{r}'] = val_date
            elif c9_sheet:
                cont_r = 7 + (idx - 7)
                if cont_r <= 50:
                    c9_sheet[f'A{cont_r}'] = self._safe_str(elig.eligibility_type)
                    c9_sheet[f'F{cont_r}'] = self._safe_str(elig.rating)
                    c9_sheet[f'G{cont_r}'] = exam_date
                    c9_sheet[f'I{cont_r}'] = self._safe_str(elig.examination_place)
                    c9_sheet[f'L{cont_r}'] = self._safe_str(elig.license_number)
                    c9_sheet[f'M{cont_r}'] = val_date

        if len(eligs) <= 7 and c9_sheet:
            c9_sheet.sheet_state = 'hidden'

        # V. WORK EXPERIENCE (rows 18 to 41 on C2: A=From, C=To, D=Position, G=Department, J=Salary, K=Salary Grade, L=Status, M=Gov)
        c6_sheet = self.wb['C6_Work Exp cont.'] if 'C6_Work Exp cont.' in self.wb.sheetnames else None
        works = list(t.work_experiences.all().order_by('-from_date', 'id'))
        for idx, work in enumerate(works):
            d_from = self._format_date(work.from_date)
            d_to = "PRESENT" if work.to_present else self._format_date(work.to_date)
            is_gov = "Y" if work.is_government_service else "N"
            salary = float(work.monthly_salary) if work.monthly_salary else ""

            if idx < 24:
                r = 18 + idx
                ws[f'A{r}'] = d_from
                ws[f'C{r}'] = d_to
                ws[f'D{r}'] = self._safe_str(work.position_title)
                ws[f'G{r}'] = self._safe_str(work.department_agency_company)
                ws[f'J{r}'] = salary
                ws[f'K{r}'] = self._safe_str(work.salary_grade)
                ws[f'L{r}'] = self._safe_str(work.appointment_status)
                ws[f'M{r}'] = is_gov
            elif c6_sheet:
                cont_r = 7 + (idx - 24)
                if cont_r <= 60:
                    c6_sheet[f'A{cont_r}'] = d_from
                    c6_sheet[f'C{cont_r}'] = d_to
                    c6_sheet[f'D{cont_r}'] = self._safe_str(work.position_title)
                    c6_sheet[f'G{cont_r}'] = self._safe_str(work.department_agency_company)
                    c6_sheet[f'J{cont_r}'] = salary
                    c6_sheet[f'K{cont_r}'] = self._safe_str(work.salary_grade)
                    c6_sheet[f'L{cont_r}'] = self._safe_str(work.appointment_status)
                    c6_sheet[f'M{cont_r}'] = is_gov

        if len(works) <= 24 and c6_sheet:
            c6_sheet.sheet_state = 'hidden'

        # Bottom signature on C2
        full_name = f"{self._safe_str(t.user.first_name if t.user else '')} {self._safe_str(t.user.last_name if t.user else '')}".strip().upper()
        ws['D43'] = full_name
        ws['I43'] = datetime.date.today().strftime("%d/%m/%Y")

    def _populate_c3(self):
        ws = self.wb['C3']
        t = self.teacher

        # VI. LEARNING AND DEVELOPMENT (L&D) (rows 5 to 21 on C3: A=Title, E=From, F=To, G=Hours, H=Type, I=Conducted By)
        c5_sheet = self.wb['C5_L&D cont.'] if 'C5_L&D cont.' in self.wb.sheetnames else None
        trainings = list(t.learning_developments.all().order_by('-from_date', 'id'))
        for idx, tr in enumerate(trainings):
            d_from = self._format_date(tr.from_date)
            d_to = self._format_date(tr.to_date)
            hrs = float(tr.number_of_hours) if tr.number_of_hours else ""
            if idx < 17:
                r = 5 + idx
                ws[f'A{r}'] = self._safe_str(tr.training_title)
                ws[f'E{r}'] = d_from
                ws[f'F{r}'] = d_to
                ws[f'G{r}'] = hrs
                ws[f'H{r}'] = self._safe_str(tr.ld_type)
                ws[f'I{r}'] = self._safe_str(tr.conducted_sponsored_by)
            elif c5_sheet:
                cont_r = 7 + (idx - 17)
                if cont_r <= 60:
                    c5_sheet[f'A{cont_r}'] = self._safe_str(tr.training_title)
                    c5_sheet[f'E{cont_r}'] = d_from
                    c5_sheet[f'F{cont_r}'] = d_to
                    c5_sheet[f'G{cont_r}'] = hrs
                    c5_sheet[f'H{cont_r}'] = self._safe_str(tr.ld_type)
                    c5_sheet[f'I{cont_r}'] = self._safe_str(tr.conducted_sponsored_by)

        if len(trainings) <= 17 and c5_sheet:
            c5_sheet.sheet_state = 'hidden'

        # VII. VOLUNTARY WORK (rows 27 to 35 on C3: A=Name & Address, E=From, F=To, G=Hours, H=Position)
        c10_sheet = self.wb['C10_Vol. Work cont.'] if 'C10_Vol. Work cont.' in self.wb.sheetnames else None
        vol_works = list(t.voluntary_works.all().order_by('-from_date', 'id'))
        for idx, vw in enumerate(vol_works):
            d_from = self._format_date(vw.from_date)
            d_to = self._format_date(vw.to_date)
            hrs = float(vw.number_of_hours) if vw.number_of_hours else ""
            if idx < 9:
                r = 27 + idx
                ws[f'A{r}'] = self._safe_str(vw.organization_name_address)
                ws[f'E{r}'] = d_from
                ws[f'F{r}'] = d_to
                ws[f'G{r}'] = hrs
                ws[f'H{r}'] = self._safe_str(vw.position_nature_of_work)
            elif c10_sheet:
                cont_r = 7 + (idx - 9)
                if cont_r <= 50:
                    c10_sheet[f'A{cont_r}'] = self._safe_str(vw.organization_name_address)
                    c10_sheet[f'E{cont_r}'] = d_from
                    c10_sheet[f'F{cont_r}'] = d_to
                    c10_sheet[f'G{cont_r}'] = hrs
                    c10_sheet[f'H{cont_r}'] = self._safe_str(vw.position_nature_of_work)

        if len(vol_works) <= 9 and c10_sheet:
            c10_sheet.sheet_state = 'hidden'

        # VIII. OTHER INFORMATION (rows 39 to 45 on C3: A=Special Skills, C=Distinctions, I=Memberships)
        c11_sheet = self.wb['C11_Other Info cont.'] if 'C11_Other Info cont.' in self.wb.sheetnames else None
        oi = getattr(t, 'other_information', None)
        skills = (oi.special_skills_hobbies if oi else []) or []
        distinctions = (oi.non_academic_distinctions if oi else []) or []
        memberships = (oi.organization_memberships if oi else []) or []

        def get_item_text(val):
            if isinstance(val, dict):
                return val.get('name') or val.get('skill') or val.get('title') or str(val)
            return str(val) if val else ""

        max_other = max(len(skills), len(distinctions), len(memberships), 0)
        for idx in range(max_other):
            sk = get_item_text(skills[idx]) if idx < len(skills) else ""
            rec = get_item_text(distinctions[idx]) if idx < len(distinctions) else ""
            org = get_item_text(memberships[idx]) if idx < len(memberships) else ""

            if idx < 7:
                r = 39 + idx
                if sk:
                    ws[f'A{r}'] = self._safe_str(sk)
                if rec:
                    ws[f'C{r}'] = self._safe_str(rec)
                if org:
                    ws[f'I{r}'] = self._safe_str(org)
            elif c11_sheet:
                cont_r = 7 + (idx - 7)
                if cont_r <= 50:
                    if sk:
                        c11_sheet[f'A{cont_r}'] = self._safe_str(sk)
                    if rec:
                        c11_sheet[f'C{cont_r}'] = self._safe_str(rec)
                    if org:
                        c11_sheet[f'I{cont_r}'] = self._safe_str(org)

        if max_other <= 7 and c11_sheet:
            c11_sheet.sheet_state = 'hidden'

        # Bottom signature on C3 (C47 is signature name, I47 is date)
        full_name = f"{self._safe_str(t.user.first_name if t.user else '')} {self._safe_str(t.user.last_name if t.user else '')}".strip().upper()
        ws['C47'] = full_name
        ws['I47'] = datetime.date.today().strftime("%d/%m/%Y")

    def _populate_c4(self):
        ws = self.wb['C4']
        t = self.teacher

        # IX. QUESTIONS 34 to 40
        q = getattr(t, 'government_questions', None)
        if q:
            # 34a
            ws['H3'] = bool(q.q34a_related_third_degree)
            ws['J3'] = not bool(q.q34a_related_third_degree)
            if q.q34a_details:
                ws['H4'] = self._safe_str(q.q34a_details)
            # 34b
            ws['H8'] = bool(q.q34b_related_fourth_degree)
            ws['J8'] = not bool(q.q34b_related_fourth_degree)
            if q.q34b_details:
                ws['H9'] = self._safe_str(q.q34b_details)

            # 35a
            ws['H13'] = bool(q.q35a_admin_offense)
            ws['J13'] = not bool(q.q35a_admin_offense)
            if q.q35a_details:
                ws['H14'] = self._safe_str(q.q35a_details)
            # 35b
            ws['H18'] = bool(q.q35b_criminal_charged)
            ws['J18'] = not bool(q.q35b_criminal_charged)
            if q.q35b_details or q.q35b_status:
                ws['H19'] = f"{self._safe_str(q.q35b_details)} {self._safe_str(q.q35b_status)}".strip()

            # 36
            ws['H23'] = bool(q.q36_convicted)
            ws['J23'] = not bool(q.q36_convicted)
            if q.q36_details:
                ws['H24'] = self._safe_str(q.q36_details)

            # 37
            ws['H27'] = bool(q.q37_separated_service)
            ws['J27'] = not bool(q.q37_separated_service)
            if q.q37_details:
                ws['H28'] = self._safe_str(q.q37_details)

            # 38a
            ws['H31'] = bool(q.q38a_candidate_election)
            ws['J31'] = not bool(q.q38a_candidate_election)
            if q.q38a_details:
                ws['H32'] = self._safe_str(q.q38a_details)
            # 38b
            ws['H34'] = bool(q.q38b_resigned_campaign)
            ws['J34'] = not bool(q.q38b_resigned_campaign)
            if q.q38b_details:
                ws['H35'] = self._safe_str(q.q38b_details)

            # 39
            ws['H37'] = bool(q.q39_immigrant_permanent_resident)
            ws['J37'] = not bool(q.q39_immigrant_permanent_resident)
            if q.q39_country:
                ws['H38'] = self._safe_str(q.q39_country)

            # 40a
            ws['H43'] = bool(q.q40a_indigenous_group)
            ws['J43'] = not bool(q.q40a_indigenous_group)
            if q.q40a_specify:
                ws['H44'] = self._safe_str(q.q40a_specify)
            # 40b
            ws['H45'] = bool(q.q40b_person_with_disability)
            ws['J45'] = not bool(q.q40b_person_with_disability)
            if q.q40b_id_number:
                ws['H46'] = self._safe_str(q.q40b_id_number)
            # 40c
            ws['H47'] = bool(q.q40c_solo_parent)
            ws['J47'] = not bool(q.q40c_solo_parent)
            if q.q40c_id_number:
                ws['H48'] = self._safe_str(q.q40c_id_number)

        # X. REFERENCES (Item 41, rows 52 to 54 on C4: A=Name, G=Address, H=Contact)
        for idx, ref in enumerate(t.references.all().order_by('id')[:3]):
            r = 52 + idx
            ws[f'A{r}'] = self._safe_str(ref.name).upper()
            ws[f'G{r}'] = self._safe_str(ref.office_residential_address)
            ws[f'H{r}'] = self._safe_str(ref.contact_no_email)

        # XI. ITEM 42 (Sworn Declaration, ID, Signature, Jurat)
        dec = getattr(t, 'pds_declaration', None)
        if dec:
            ws['D61'] = self._safe_str(dec.government_id_type)
            ws['D62'] = self._safe_str(dec.id_license_passport_no)
            ws['D64'] = self._safe_str(dec.date_place_of_issuance)
            ws['G64'] = self._format_date(dec.date_accomplished or datetime.date.today())

            if dec.administering_officer_name:
                officer_txt = self._safe_str(dec.administering_officer_name).upper()
                if dec.administering_officer_title:
                    officer_txt += f"\n{self._safe_str(dec.administering_officer_title)}"
                ws['F69'] = officer_txt
            if dec.subscribed_date:
                ws['A67'] = f"SUBSCRIBED AND SWORN to before me this {dec.subscribed_date}, affiant exhibiting his/her validly issued government ID as indicated above."
        else:
            ws['G64'] = datetime.date.today().strftime("%d/%m/%Y")

        # Signature printed name inside signature box G60
        full_name = f"{self._safe_str(t.user.first_name if t.user else '')} {self._safe_str(t.user.last_name if t.user else '')}".strip().upper()
        ws['G60'] = full_name

        # Embed Signature Image if available
        if t.signature_image and hasattr(t.signature_image, 'path') and os.path.exists(t.signature_image.path):
            try:
                sig_img = OpenpyxlImage(t.signature_image.path)
                sig_img.width = 160
                sig_img.height = 45
                ws.add_image(sig_img, 'G63')
            except Exception as e:
                print("Failed to embed signature into Excel:", e)

        # Embed Photo Image if available
        if t.photo and hasattr(t.photo, 'path') and os.path.exists(t.photo.path):
            try:
                photo_img = OpenpyxlImage(t.photo.path)
                photo_img.width = 135
                photo_img.height = 165
                ws.add_image(photo_img, 'L56')
            except Exception as e:
                print("Failed to embed photo into Excel:", e)
