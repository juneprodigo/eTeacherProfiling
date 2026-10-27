import io
from datetime import date
from django.http import HttpResponse
from openpyxl import Workbook
from openpyxl.styles import Font, Alignment, PatternFill, Border, Side
from openpyxl.utils import get_column_letter
from .models import StepIncrementRecord, ServiceRecordEntry, SalaryAdjustmentRecord
from pds.models import Teacher


def generate_deped_nosi_excel(record_id):
    """
    Generates a professionally styled Official DepEd Form: Notice of Step Increment (NOSI)
    conforming with standard Civil Service Commission & Department of Education format.
    """
    record = StepIncrementRecord.objects.select_related('teacher__user').get(id=record_id)
    wb = Workbook()
    ws = wb.active
    ws.title = "DepEd NOSI"

    # Ensure grid lines are visible
    ws.views.sheetView[0].showGridLines = True

    # Styling constants
    font_header = Font(name="Calibri", size=11, bold=True)
    font_bold = Font(name="Calibri", size=10, bold=True)
    font_regular = Font(name="Calibri", size=10)
    font_italic = Font(name="Calibri", size=9, italic=True)
    font_title = Font(name="Arial", size=13, bold=True)
    align_center = Alignment(horizontal="center", vertical="center", wrap_text=True)
    align_left = Alignment(horizontal="left", vertical="center", wrap_text=True)
    align_right = Alignment(horizontal="right", vertical="center")

    thin_border_side = Side(border_style="thin", color="000000")
    double_bottom_side = Side(border_style="double", color="000000")
    box_border = Border(left=thin_border_side, right=thin_border_side, top=thin_border_side, bottom=thin_border_side)
    summary_border = Border(top=thin_border_side, bottom=double_bottom_side)

    # Column Widths
    ws.column_dimensions['A'].width = 6
    ws.column_dimensions['B'].width = 24
    ws.column_dimensions['C'].width = 20
    ws.column_dimensions['D'].width = 20
    ws.column_dimensions['E'].width = 20
    ws.column_dimensions['F'].width = 18

    # Header section
    ws.merge_cells('A2:F2')
    ws['A2'] = "Republic of the Philippines"
    ws['A2'].alignment = align_center
    ws['A2'].font = font_regular

    ws.merge_cells('A3:F3')
    ws['A3'] = "Department of Education"
    ws['A3'].alignment = align_center
    ws['A3'].font = font_header

    ws.merge_cells('A4:F4')
    ws['A4'] = record.division_office or "Schools Division Office"
    ws['A4'].alignment = align_center
    ws['A4'].font = font_regular

    ws.merge_cells('A5:F5')
    ws['A5'] = record.school_name or "District / School"
    ws['A5'].alignment = align_center
    ws['A5'].font = font_regular

    # Title
    ws.merge_cells('A7:F7')
    ws['A7'] = "NOTICE OF STEP INCREMENT"
    ws['A7'].alignment = align_center
    ws['A7'].font = font_title

    # NOSI metadata
    ws['A9'] = "Date:"
    ws['A9'].font = font_bold
    ws['B9'] = record.date_issued.strftime('%B %d, %Y') if record.date_issued else date.today().strftime('%B %d, %Y')
    ws['B9'].font = font_regular

    ws['E9'] = "NOSI No.:"
    ws['E9'].font = font_bold
    ws['F9'] = record.nosi_number or f"NOSI-{record.id:04d}"
    ws['F9'].font = font_bold

    # Addressee
    teacher_name = record.teacher.user.get_full_name().upper() or record.teacher.user.username.upper()
    ws['A11'] = "TO:"
    ws['A11'].font = font_bold
    ws['B11'] = teacher_name
    ws['B11'].font = font_header

    ws['B12'] = record.teacher.position or "Teacher"
    ws['B12'].font = font_regular
    ws['B13'] = record.school_name or "Department of Education"
    ws['B13'].font = font_regular

    # Body Paragraph
    ws.merge_cells('A15:F17')
    body_text = (
        f"Pursuant to the provisions of Joint CSC-DBM Circular No. 1, s. 2012 / DepEd Guidelines, your salary as "
        f"{record.teacher.position or 'Teacher'} is hereby adjusted effective {record.effective_date.strftime('%B %d, %Y')}, "
        f"as an entitlement to Step Increment based on {record.get_basis_display()}."
    )
    ws['A15'] = body_text
    ws['A15'].alignment = align_left
    ws['A15'].font = font_regular

    # Table breakdown
    ws.merge_cells('A19:D19')
    ws['A19'] = "PARTICULARS"
    ws['A19'].font = font_bold
    ws['A19'].alignment = align_left
    ws['A19'].fill = PatternFill(start_color="F2F4F8", end_color="F2F4F8", fill_type="solid")
    ws.merge_cells('E19:F19')
    ws['E19'] = "SALARY DETAILS"
    ws['E19'].font = font_bold
    ws['E19'].alignment = align_center
    ws['E19'].fill = PatternFill(start_color="F2F4F8", end_color="F2F4F8", fill_type="solid")

    for col in ['A', 'B', 'C', 'D', 'E', 'F']:
        ws[f'{col}19'].border = box_border

    # Previous Salary
    ws.merge_cells('A20:D20')
    ws['A20'] = f"1. Actual Basic Salary Prior to Increment (SG {record.previous_salary_grade}, Step {record.previous_step})"
    ws['A20'].font = font_regular
    ws.merge_cells('E20:F20')
    ws['E20'] = f"₱ {record.previous_monthly_salary:,.2f}"
    ws['E20'].font = font_regular
    ws['E20'].alignment = align_right

    # New Salary
    ws.merge_cells('A21:D21')
    ws['A21'] = f"2. Adjusted Basic Salary per NOSI (SG {record.new_salary_grade}, Step {record.new_step})"
    ws['A21'].font = font_regular
    ws.merge_cells('E21:F21')
    ws['E21'] = f"₱ {record.new_monthly_salary:,.2f}"
    ws['E21'].font = font_bold
    ws['E21'].alignment = align_right

    # Monthly Differential
    ws.merge_cells('A22:D22')
    ws['A22'] = "3. Monthly Salary Differential (Add-on to basic pay)"
    ws['A22'].font = font_bold
    ws.merge_cells('E22:F22')
    ws['E22'] = f"₱ {record.monthly_differential:,.2f}"
    ws['E22'].font = font_bold
    ws['E22'].alignment = align_right

    for r in range(20, 23):
        for c in range(1, 7):
            cell = ws.cell(row=r, column=c)
            cell.border = box_border

    # Condition notes
    ws.merge_cells('A24:F26')
    condition_text = (
        "This step increment is subject to post-audit by the Commission on Audit (COA) and to the condition that you "
        "have rendered continuous, satisfactory service. Any excess payment made shall be immediately refunded to DepEd."
    )
    ws['A24'] = condition_text
    ws['A24'].font = font_italic
    ws['A24'].alignment = align_left

    # Checklist block
    ws.merge_cells('A28:F28')
    ws['A28'] = "PROCESSING CHECKLIST & KEY REQUIREMENTS (DepEd Differential Claim)"
    ws['A28'].font = font_bold
    ws['A28'].fill = PatternFill(start_color="EBF3FA", end_color="EBF3FA", fill_type="solid")
    ws['A28'].alignment = align_left

    reqs = [
        ("Approved NOSI", record.has_approved_nosi),
        ("Annex D / Relevant forms", record.has_annex_d),
        ("Payslips or payroll records (before and after step inclusion)", record.has_payslips),
        ("Latest Appointment", record.has_latest_appointment),
        ("Updated Service Record", record.has_updated_service_record),
    ]

    for idx, (label, is_checked) in enumerate(reqs, start=29):
        check_symbol = "[ ✓ ]" if is_checked else "[   ]"
        status_text = "Verified / Attached" if is_checked else "Pending Requirement"
        ws.merge_cells(f'A{idx}:D{idx}')
        ws[f'A{idx}'] = f"{check_symbol}  {label}"
        ws[f'A{idx}'].font = font_bold if is_checked else font_regular
        ws.merge_cells(f'E{idx}:F{idx}')
        ws[f'E{idx}'] = status_text
        ws[f'E{idx}'].font = font_italic
        ws[f'E{idx}'].alignment = align_right

        for c in range(1, 7):
            ws.cell(row=idx, column=c).border = box_border

    # Signatures
    sign_row = 36
    ws.merge_cells(f'A{sign_row}:C{sign_row}')
    ws[f'A{sign_row}'] = "Prepared by / HRMO:"
    ws[f'A{sign_row}'].font = font_regular

    ws.merge_cells(f'D{sign_row}:F{sign_row}')
    ws[f'D{sign_row}'] = "Approved by:"
    ws[f'D{sign_row}'].font = font_regular

    name_row = sign_row + 3
    ws.merge_cells(f'A{name_row}:C{name_row}')
    ws[f'A{name_row}'] = "ADMINISTRATIVE OFFICER / HRMO"
    ws[f'A{name_row}'].font = font_bold
    ws[f'A{name_row}'].alignment = align_center

    ws.merge_cells(f'D{name_row}:F{name_row}')
    ws[f'D{name_row}'] = (record.schools_division_superintendent or "SCHOOLS DIVISION SUPERINTENDENT").upper()
    ws[f'D{name_row}'].font = font_bold
    ws[f'D{name_row}'].alignment = align_center

    title_row = name_row + 1
    ws.merge_cells(f'D{title_row}:F{title_row}')
    ws[f'D{title_row}'] = "Schools Division Superintendent / Authorized Official"
    ws[f'D{title_row}'].font = font_italic
    ws[f'D{title_row}'].alignment = align_center

    output = io.BytesIO()
    wb.save(output)
    output.seek(0)
    return output.getvalue()


def generate_csc_service_record_excel(teacher_id):
    """
    Generates an official Civil Service Commission Form No. 122 (Service Record)
    as used in DepEd and the Philippine Civil Service.
    """
    teacher = Teacher.objects.select_related('user').get(id=teacher_id)
    records = ServiceRecordEntry.objects.filter(teacher=teacher).order_by('from_date', 'id')

    wb = Workbook()
    ws = wb.active
    ws.title = "CSC Form 122 - Service Record"
    ws.views.sheetView[0].showGridLines = True

    font_header = Font(name="Calibri", size=11, bold=True)
    font_bold = Font(name="Calibri", size=9, bold=True)
    font_regular = Font(name="Calibri", size=9)
    font_italic = Font(name="Calibri", size=8, italic=True)
    font_title = Font(name="Arial", size=12, bold=True)
    align_center = Alignment(horizontal="center", vertical="center", wrap_text=True)
    align_left = Alignment(horizontal="left", vertical="center", wrap_text=True)
    align_right = Alignment(horizontal="right", vertical="center")

    thin_border_side = Side(border_style="thin", color="000000")
    box_border = Border(left=thin_border_side, right=thin_border_side, top=thin_border_side, bottom=thin_border_side)
    header_fill = PatternFill(start_color="F1F5F9", end_color="F1F5F9", fill_type="solid")

    # Column Widths
    col_widths = {
        'A': 13, # From Date
        'B': 13, # To Date
        'C': 26, # Designation
        'D': 14, # Status
        'E': 14, # Salary (Annual/Monthly)
        'F': 24, # Station / Place
        'G': 12, # Branch
        'H': 14, # LWOP
        'I': 13, # Sep Date
        'J': 16, # Sep Cause
        'K': 20, # Remarks
    }
    for col, w in col_widths.items():
        ws.column_dimensions[col].width = w

    # Header section
    ws.merge_cells('A1:K1')
    ws['A1'] = "Republic of the Philippines"
    ws['A1'].alignment = align_center
    ws['A1'].font = font_regular

    ws.merge_cells('A2:K2')
    ws['A2'] = "DEPARTMENT OF EDUCATION"
    ws['A2'].alignment = align_center
    ws['A2'].font = font_header

    ws.merge_cells('A3:K3')
    ws['A3'] = "Region II - Cagayan Valley • Schools Division of Isabela"
    ws['A3'].alignment = align_center
    ws['A3'].font = font_regular

    ws.merge_cells('A5:K5')
    ws['A5'] = "SERVICE RECORD"
    ws['A5'].alignment = align_center
    ws['A5'].font = font_title

    ws.merge_cells('A6:K6')
    ws['A6'] = "(Civil Service Commission Form No. 122)"
    ws['A6'].alignment = align_center
    ws['A6'].font = font_italic

    # Teacher Info
    personal_info = getattr(teacher, 'personal_info', None)
    dob_str = personal_info.date_of_birth.strftime('%B %d, %Y') if (personal_info and personal_info.date_of_birth) else 'N/A'
    pob_str = personal_info.place_of_birth if personal_info else 'N/A'
    emp_no = teacher.employee_id or 'N/A'
    full_name = teacher.user.get_full_name().upper() or teacher.user.username.upper()

    ws['A8'] = "NAME:"
    ws['A8'].font = font_bold
    ws['B8'] = full_name
    ws['B8'].font = font_header

    ws['F8'] = "BIRTH DATE:"
    ws['F8'].font = font_bold
    ws['G8'] = dob_str
    ws['G8'].font = font_regular

    ws['I8'] = "BIRTH PLACE:"
    ws['I8'].font = font_bold
    ws['J8'] = pob_str
    ws['J8'].font = font_regular

    ws['A9'] = "EMPLOYEE NO.:"
    ws['A9'].font = font_bold
    ws['B9'] = emp_no
    ws['B9'].font = font_bold

    ws['F9'] = "CURRENT POSITION:"
    ws['F9'].font = font_bold
    ws['G9'] = teacher.position or "Teacher"
    ws['G9'].font = font_regular

    # Table Headers
    # Row 11: Top grouped headers
    ws.merge_cells('A11:E11')
    ws['A11'] = "SERVICE RECORD / APPOINTMENT HISTORY"
    ws['A11'].alignment = align_center
    ws['A11'].font = font_bold
    ws['A11'].fill = header_fill

    ws.merge_cells('F11:G11')
    ws['F11'] = "OFFICE ENTITY / DIV"
    ws['F11'].alignment = align_center
    ws['F11'].font = font_bold
    ws['F11'].fill = header_fill

    ws.merge_cells('H11:J11')
    ws['H11'] = "SEPARATION / ABSENCE"
    ws['H11'].alignment = align_center
    ws['H11'].font = font_bold
    ws['H11'].fill = header_fill

    ws['K11'] = "REMARKS"
    ws['K11'].alignment = align_center
    ws['K11'].font = font_bold
    ws['K11'].fill = header_fill

    # Row 12: Column headers
    headers_r12 = [
        ('A', 'From Date'),
        ('B', 'To Date'),
        ('C', 'Designation / Rank'),
        ('D', 'Status'),
        ('E', 'Monthly / Annual Salary'),
        ('F', 'Station / School'),
        ('G', 'Branch'),
        ('H', 'L/V WOP Dates'),
        ('I', 'Sep Date'),
        ('J', 'Cause'),
        ('K', 'Remarks'),
    ]
    for col, text in headers_r12:
        cell = ws[f'{col}12']
        cell.value = text
        cell.alignment = align_center
        cell.font = font_bold
        cell.fill = header_fill

    for r in range(11, 13):
        for c in range(1, 12):
            ws.cell(row=r, column=c).border = box_border

    # Table Rows
    curr_row = 13
    if not records.exists():
        ws.merge_cells(f'A{curr_row}:K{curr_row}')
        ws[f'A{curr_row}'] = "No official service record entries encoded yet."
        ws[f'A{curr_row}'].alignment = align_center
        ws[f'A{curr_row}'].font = font_italic
        for c in range(1, 12):
            ws.cell(row=curr_row, column=c).border = box_border
        curr_row += 1
    else:
        for rec in records:
            from_str = rec.from_date.strftime('%m/%d/%Y') if rec.from_date else ''
            to_str = "Present" if rec.to_present else (rec.to_date.strftime('%m/%d/%Y') if rec.to_date else '')
            sal_str = f"₱{rec.monthly_salary:,.2f}" if rec.monthly_salary else (f"₱{rec.annual_salary:,.2f}/yr" if rec.annual_salary else '')
            sep_date_str = rec.separation_date.strftime('%m/%d/%Y') if rec.separation_date else ''

            ws[f'A{curr_row}'] = from_str
            ws[f'A{curr_row}'].alignment = align_center
            ws[f'A{curr_row}'].font = font_regular

            ws[f'B{curr_row}'] = to_str
            ws[f'B{curr_row}'].alignment = align_center
            ws[f'B{curr_row}'].font = font_regular

            ws[f'C{curr_row}'] = rec.designation
            ws[f'C{curr_row}'].alignment = align_left
            ws[f'C{curr_row}'].font = font_regular

            ws[f'D{curr_row}'] = rec.status_of_appointment
            ws[f'D{curr_row}'].alignment = align_center
            ws[f'D{curr_row}'].font = font_regular

            ws[f'E{curr_row}'] = sal_str
            ws[f'E{curr_row}'].alignment = align_right
            ws[f'E{curr_row}'].font = font_regular

            ws[f'F{curr_row}'] = rec.station_office
            ws[f'F{curr_row}'].alignment = align_left
            ws[f'F{curr_row}'].font = font_regular

            ws[f'G{curr_row}'] = rec.branch
            ws[f'G{curr_row}'].alignment = align_center
            ws[f'G{curr_row}'].font = font_regular

            ws[f'H{curr_row}'] = rec.lwop_dates or "NONE"
            ws[f'H{curr_row}'].alignment = align_center
            ws[f'H{curr_row}'].font = font_regular

            ws[f'I{curr_row}'] = sep_date_str
            ws[f'I{curr_row}'].alignment = align_center
            ws[f'I{curr_row}'].font = font_regular

            ws[f'J{curr_row}'] = rec.separation_cause or ''
            ws[f'J{curr_row}'].alignment = align_left
            ws[f'J{curr_row}'].font = font_regular

            ws[f'K{curr_row}'] = rec.remarks or ''
            ws[f'K{curr_row}'].alignment = align_left
            ws[f'K{curr_row}'].font = font_regular

            for c in range(1, 12):
                ws.cell(row=curr_row, column=c).border = box_border

            curr_row += 1

    # Certification Block
    curr_row += 2
    ws.merge_cells(f'A{curr_row}:K{curr_row}')
    cert_text = (
        "Issued in compliance with Executive Order No. 54 dated August 10, 1954, and in accordance with "
        "Circular No. 58, dated August 10, 1954 of the System."
    )
    ws[f'A{curr_row}'] = cert_text
    ws[f'A{curr_row}'].alignment = align_left
    ws[f'A{curr_row}'].font = font_italic

    curr_row += 3
    ws['B{0}'.format(curr_row)] = f"Date: {date.today().strftime('%B %d, %Y')}"
    ws['B{0}'.format(curr_row)].font = font_regular

    ws.merge_cells(f'H{curr_row}:K{curr_row}')
    ws[f'H{curr_row}'] = "CERTIFIED CORRECT:"
    ws[f'H{curr_row}'].font = font_bold
    ws[f'H{curr_row}'].alignment = align_center

    curr_row += 3
    ws.merge_cells(f'H{curr_row}:K{curr_row}')
    ws[f'H{curr_row}'] = "ADMINISTRATIVE OFFICER V / HRMO"
    ws[f'H{curr_row}'].font = font_bold
    ws[f'H{curr_row}'].alignment = align_center

    curr_row += 1
    ws.merge_cells(f'H{curr_row}:K{curr_row}')
    ws[f'H{curr_row}'] = "Schools Division Office of Isabela"
    ws[f'H{curr_row}'].font = font_italic
    ws[f'H{curr_row}'].alignment = align_center

    output = io.BytesIO()
    wb.save(output)
    output.seek(0)
    return output.getvalue()


def generate_deped_nosa_excel(nosa_id):
    """
    Generates an official DepEd Notice of Salary Adjustment (NOSA) spreadsheet
    conforming with standard National Budget Circular / SSL format.
    """
    nosa = SalaryAdjustmentRecord.objects.select_related('teacher__user').get(id=nosa_id)
    wb = Workbook()
    ws = wb.active
    ws.title = "DepEd NOSA"
    ws.views.sheetView[0].showGridLines = True

    font_header = Font(name="Calibri", size=11, bold=True)
    font_bold = Font(name="Calibri", size=10, bold=True)
    font_regular = Font(name="Calibri", size=10)
    font_italic = Font(name="Calibri", size=9, italic=True)
    font_title = Font(name="Arial", size=13, bold=True)
    align_center = Alignment(horizontal="center", vertical="center", wrap_text=True)
    align_left = Alignment(horizontal="left", vertical="center", wrap_text=True)
    align_right = Alignment(horizontal="right", vertical="center")

    thin_border_side = Side(border_style="thin", color="000000")
    double_bottom_side = Side(border_style="double", color="000000")
    box_border = Border(left=thin_border_side, right=thin_border_side, top=thin_border_side, bottom=thin_border_side)
    summary_border = Border(top=thin_border_side, bottom=double_bottom_side)

    ws.column_dimensions['A'].width = 8
    ws.column_dimensions['B'].width = 24
    ws.column_dimensions['C'].width = 20
    ws.column_dimensions['D'].width = 20
    ws.column_dimensions['E'].width = 20
    ws.column_dimensions['F'].width = 18

    # Header
    ws.merge_cells('A2:F2')
    ws['A2'] = "Republic of the Philippines"
    ws['A2'].alignment = align_center
    ws['A2'].font = font_regular

    ws.merge_cells('A3:F3')
    ws['A3'] = "Department of Education"
    ws['A3'].alignment = align_center
    ws['A3'].font = font_header

    ws.merge_cells('A4:F4')
    ws['A4'] = "Schools Division Office of Isabela"
    ws['A4'].alignment = align_center
    ws['A4'].font = font_regular

    ws.merge_cells('A7:F7')
    ws['A7'] = "NOTICE OF SALARY ADJUSTMENT (NOSA)"
    ws['A7'].alignment = align_center
    ws['A7'].font = font_title

    ws['A9'] = "Date:"
    ws['A9'].font = font_bold
    ws['B9'] = nosa.date_issued.strftime('%B %d, %Y') if nosa.date_issued else date.today().strftime('%B %d, %Y')
    ws['B9'].font = font_regular

    ws['E9'] = "NOSA No.:"
    ws['E9'].font = font_bold
    ws['F9'] = nosa.nosa_number or f"NOSA-{nosa.id:04d}"
    ws['F9'].font = font_bold

    teacher_name = nosa.teacher.user.get_full_name().upper() or nosa.teacher.user.username.upper()
    ws['A11'] = "TO:"
    ws['A11'].font = font_bold
    ws['B11'] = teacher_name
    ws['B11'].font = font_header

    ws['B12'] = nosa.position_title or nosa.teacher.position or "Teacher"
    ws['B12'].font = font_regular

    ws.merge_cells('A15:F17')
    body_text = (
        f"Pursuant to {nosa.legal_basis}, your salary as {nosa.position_title or 'Teacher'} "
        f"is hereby adjusted effective {nosa.effective_date.strftime('%B %d, %Y')}, as follows:"
    )
    ws['A15'] = body_text
    ws['A15'].alignment = align_left
    ws['A15'].font = font_regular

    # Particulars Table
    ws.merge_cells('A19:D19')
    ws['A19'] = "PARTICULARS"
    ws['A19'].font = font_bold
    ws.merge_cells('E19:F19')
    ws['E19'] = "AMOUNT (PHP)"
    ws['E19'].font = font_bold
    ws['E19'].alignment = align_right

    ws.merge_cells('A20:D20')
    ws['A20'] = f"1. Adjusted Monthly Basic Salary (SG {nosa.salary_grade}, Step {nosa.step}):"
    ws['A20'].font = font_regular
    ws.merge_cells('E20:F20')
    ws['E20'] = f"₱{nosa.adjusted_monthly_salary:,.2f}"
    ws['E20'].font = font_bold
    ws['E20'].alignment = align_right

    ws.merge_cells('A21:D21')
    ws['A21'] = f"2. Actual Monthly Basic Salary Prior to Adjustment:"
    ws['A21'].font = font_regular
    ws.merge_cells('E21:F21')
    ws['E21'] = f"₱{nosa.previous_monthly_salary:,.2f}"
    ws['E21'].font = font_regular
    ws['E21'].alignment = align_right

    ws.merge_cells('A22:D22')
    ws['A22'] = "3. Monthly Salary Differential:"
    ws['A22'].font = font_bold
    ws.merge_cells('E22:F22')
    ws['E22'] = f"₱{nosa.monthly_differential:,.2f}"
    ws['E22'].font = font_bold
    ws['E22'].alignment = align_right

    for r in range(19, 23):
        for c in range(1, 7):
            ws.cell(row=r, column=c).border = box_border

    ws.merge_cells('A25:F26')
    ws['A25'] = (
        "This salary adjustment is subject to post-audit, and to prompt refund of any excess payments "
        "if found not in order."
    )
    ws['A25'].font = font_italic
    ws['A25'].alignment = align_left

    # Signatures
    sign_row = 30
    ws.merge_cells(f'D{sign_row}:F{sign_row}')
    ws[f'D{sign_row}'] = "Very truly yours,"
    ws[f'D{sign_row}'].font = font_regular

    name_row = sign_row + 3
    ws.merge_cells(f'D{name_row}:F{name_row}')
    ws[f'D{name_row}'] = (nosa.signatory_authority or "SCHOOLS DIVISION SUPERINTENDENT").upper()
    ws[f'D{name_row}'].font = font_bold
    ws[f'D{name_row}'].alignment = align_center

    title_row = name_row + 1
    ws.merge_cells(f'D{title_row}:F{title_row}')
    ws[f'D{title_row}'] = "Schools Division Superintendent"
    ws[f'D{title_row}'].font = font_italic
    ws[f'D{title_row}'].alignment = align_center

    output = io.BytesIO()
    wb.save(output)
    output.seek(0)
    return output.getvalue()

